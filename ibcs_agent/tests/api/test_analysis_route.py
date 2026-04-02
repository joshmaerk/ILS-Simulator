"""Tests for POST /v1/analyze-file."""

from __future__ import annotations

import base64
import io

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation


def _make_pptx_b64() -> str:
    """Create a minimal in-memory PPTX and return base64-encoded string."""
    prs = Presentation()
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    if title:
        title.text = "Test Slide"
    buf = io.BytesIO()
    prs.save(buf)
    return base64.b64encode(buf.getvalue()).decode()


def _make_xlsx_b64() -> str:
    """Create a minimal in-memory XLSX and return base64-encoded string."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "AC"
    ws["B1"] = "PL"
    ws["A2"] = 100
    ws["B2"] = 90
    buf = io.BytesIO()
    wb.save(buf)
    return base64.b64encode(buf.getvalue()).decode()


def test_analyze_valid_pptx(client: TestClient) -> None:
    payload = {
        "file_content_base64": _make_pptx_b64(),
        "file_name": "test.pptx",
    }
    resp = client.post("/v1/analyze-file", json=payload)
    assert resp.status_code == 200
    data = resp.json()
    assert "overall_score" in data
    assert "overall_grade" in data
    assert data["overall_grade"] in ("A", "B", "C", "D", "F")
    assert "pages" in data
    assert "total_violations" in data


def test_analyze_valid_xlsx(client: TestClient) -> None:
    payload = {
        "file_content_base64": _make_xlsx_b64(),
        "file_name": "test.xlsx",
    }
    resp = client.post("/v1/analyze-file", json=payload)
    assert resp.status_code == 200
    data = resp.json()
    assert "overall_score" in data


def test_analyze_unsupported_extension(client: TestClient) -> None:
    payload = {
        "file_content_base64": base64.b64encode(b"dummy").decode(),
        "file_name": "test.docx",
    }
    resp = client.post("/v1/analyze-file", json=payload)
    # Unsupported type → 400 or 500 with error
    assert resp.status_code in (400, 500)
    assert "error" in resp.json() or "detail" in resp.json()


def test_analyze_invalid_base64(client: TestClient) -> None:
    payload = {
        "file_content_base64": "not-valid-base64!!!",
        "file_name": "test.pptx",
    }
    resp = client.post("/v1/analyze-file", json=payload)
    assert resp.status_code in (400, 422, 500)


def test_analyze_missing_filename(client: TestClient) -> None:
    payload = {"file_content_base64": base64.b64encode(b"dummy").decode()}
    resp = client.post("/v1/analyze-file", json=payload)
    assert resp.status_code == 422  # Pydantic validation error


def test_analyze_empty_body(client: TestClient) -> None:
    resp = client.post("/v1/analyze-file", json={})
    assert resp.status_code == 422


def test_analyze_response_has_correlation_id_header(client: TestClient) -> None:
    payload = {
        "file_content_base64": _make_pptx_b64(),
        "file_name": "test.pptx",
    }
    resp = client.post("/v1/analyze-file", json=payload)
    assert "x-correlation-id" in {k.lower() for k in resp.headers}


def test_analyze_correlation_id_echoed(client: TestClient) -> None:
    cid = "test-correlation-id-12345"
    payload = {
        "file_content_base64": _make_pptx_b64(),
        "file_name": "test.pptx",
    }
    resp = client.post("/v1/analyze-file", json=payload, headers={"X-Correlation-ID": cid})
    assert resp.headers.get("x-correlation-id") == cid
